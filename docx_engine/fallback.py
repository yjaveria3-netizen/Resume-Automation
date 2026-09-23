import shutil
from contextlib import contextmanager
from pathlib import Path
import docx
from docx_engine.parser import SUBSEQUENT_HEADING_REGEX


@contextmanager
def with_document_backup(doc_path: str):
    """Context manager that creates a temporary backup before editing and rolls back if an exception occurs."""
    path = Path(doc_path)
    if not path.exists():
        yield path
        return

    backup_path = path.parent / f"{path.stem}_backup{path.suffix}"
    shutil.copy2(path, backup_path)
    try:
        yield path
        if backup_path.exists():
            backup_path.unlink()
    except Exception as e:
        if backup_path.exists():
            shutil.copy2(backup_path, path)
            backup_path.unlink()
        raise e


def create_projects_section_fallback(doc) -> int:
    """Scans doc to find best insertion spot (before Education/Skills or at end), inserts a new Projects heading, and returns its index."""
    best_idx = len(doc.paragraphs)

    # Search for Education or Skills heading to insert before
    for idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue
        if SUBSEQUENT_HEADING_REGEX.search(text):
            style_name = (p.style.name if p.style else "").lower()
            if "heading" in style_name or "title" in style_name or any(r.bold for r in p.runs if r.bold is not None):
                best_idx = idx
                break

    sample_heading_run = None
    for p in doc.paragraphs:
        if p.text.strip() and ("heading" in (p.style.name or "").lower() or any(r.bold for r in p.runs if r.bold is not None)):
            if p.runs:
                sample_heading_run = p.runs[0]
            break

    # Insert new Projects heading
    if best_idx < len(doc.paragraphs):
        target_p = doc.paragraphs[best_idx]
        new_heading_p = target_p.insert_paragraph_before("Projects")
        target_p._p.addprevious(new_heading_p._p)
    else:
        new_heading_p = doc.add_paragraph("Projects")

    run = new_heading_p.runs[0] if new_heading_p.runs else new_heading_p.add_run("Projects")
    run.bold = True
    if sample_heading_run and sample_heading_run.font.name:
        run.font.name = sample_heading_run.font.name
        run.font.size = sample_heading_run.font.size or docx.shared.Pt(14)
    else:
        run.font.name = "Calibri"
        run.font.size = docx.shared.Pt(14)

    new_heading_p.paragraph_format.space_before = docx.shared.Pt(12)
    new_heading_p.paragraph_format.space_after = docx.shared.Pt(6)

    for idx, p in enumerate(doc.paragraphs):
        if p._p == new_heading_p._p:
            return idx

    return len(doc.paragraphs) - 1
