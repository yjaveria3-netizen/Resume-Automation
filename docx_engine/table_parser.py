import docx
from typing import Optional, Tuple
from docx_engine.parser import PROJECTS_HEADING_REGEX
from docx_engine.formatter import clone_font_style, clone_paragraph_style


def find_projects_in_tables(doc) -> Optional[Tuple[int, int, int]]:
    """Scans doc.tables for cells containing Projects heading. Returns (table_idx, row_idx, col_idx) or None."""
    if not doc or not hasattr(doc, "tables"):
        return None

    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    text = p.text.strip()
                    if text and PROJECTS_HEADING_REGEX.search(text):
                        return (t_idx, r_idx, c_idx)
    return None


def insert_project_into_cell(
    cell,
    project_title: str,
    tech_stack: list,
    bullet_point: str,
) -> bool:
    """Inserts project title and bullet point inside a Word table cell."""
    tech_stack_str = ", ".join(tech_stack) if tech_stack else ""
    header_text = f"{project_title} | {tech_stack_str}" if tech_stack_str else project_title

    # Sample existing paragraph in cell if available
    sample_p = cell.paragraphs[0] if cell.paragraphs else None

    # Title paragraph inside cell
    title_p = cell.add_paragraph()
    run_title = title_p.add_run(header_text)
    run_title.bold = True
    if sample_p:
        clone_paragraph_style(sample_p, title_p)
        if sample_p.runs:
            clone_font_style(sample_p.runs[0], run_title)

    # Bullet paragraph inside cell
    clean_bullet = bullet_point.lstrip("-*• ").strip()
    bullet_p = cell.add_paragraph()
    try:
        bullet_p.style = "List Bullet"
    except Exception:
        pass
    run_bullet = bullet_p.add_run(f"• {clean_bullet}" if not bullet_p.style else clean_bullet)
    if sample_p:
        clone_paragraph_style(sample_p, bullet_p)
        if sample_p.runs:
            clone_font_style(sample_p.runs[0], run_bullet)

    return True
