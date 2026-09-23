import sys
from pathlib import Path
import docx

ROOT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT_DIR))

from docx_engine.unified_updater import update_resume_with_project


def create_no_projects_resume(path):
    doc = docx.Document()
    doc.add_heading("Javeria Yasin", 0)
    doc.add_heading("Work Experience", 1)
    doc.add_paragraph("Software Intern at Tech Corp")
    doc.add_heading("Education", 1)
    doc.add_paragraph("BS Computer Science")
    doc.save(path)


def create_table_resume(path):
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=2)
    cell_left = table.cell(0, 0)
    cell_right = table.cell(0, 1)

    p_left = cell_left.paragraphs[0]
    p_left.add_run("Projects").bold = True
    p_left_desc = cell_left.add_paragraph("Existing table project entry")

    p_right = cell_right.paragraphs[0]
    p_right.add_run("Education").bold = True

    doc.save(path)


def create_custom_headings_resume(path):
    doc = docx.Document()
    doc.add_paragraph("PORTFOLIO HIGHLIGHTS").style = "Heading 1"
    doc.add_paragraph("• Existing portfolio project entry")
    doc.save(path)


def create_custom_bullets_resume(path):
    doc = docx.Document()
    p_h = doc.add_paragraph()
    p_h.add_run("Key Projects").bold = True
    p_b = doc.add_paragraph("➤ Custom arrow bullet project entry")
    doc.save(path)


def create_sparse_doc(path):
    doc = docx.Document()
    doc.add_paragraph("Javeria Yasin - Developer Resume")
    doc.save(path)


def main():
    exp_dir = Path(__file__).resolve().parent
    print("=================================================================================")
    print("   PERSON D: DAY 5 EDGE CASES TEST HARNESS")
    print("=================================================================================")

    cases = [
        ("No Projects Section", create_no_projects_resume, "edge_no_projects.docx"),
        ("Table-Based Resume Layout", create_table_resume, "edge_table_layout.docx"),
        ("Non-Standard Headings", create_custom_headings_resume, "edge_custom_headings.docx"),
        ("Custom Bullet Glyphs", create_custom_bullets_resume, "edge_custom_bullets.docx"),
        ("Sparse 1-Paragraph Resume", create_sparse_doc, "edge_sparse_doc.docx"),
    ]

    for name, creator_fn, filename in cases:
        input_file = exp_dir / filename
        output_file = exp_dir / f"output_{filename}"

        # 1. Generate edge case document
        creator_fn(str(input_file))

        # 2. Run unified update engine
        res = update_resume_with_project(
            input_docx_path=str(input_file),
            output_docx_path=str(output_file),
            project_title="Edge Case Tested Project",
            tech_stack=["Python", "FastAPI", "Docker"],
            bullet_point="Engineered robust fallback handling using Python and FastAPI to ensure 100% document update reliability."
        )

        print(f"[{name}]")
        print(f"  Result: Success={res['success']}, Method={res['method']}, Output={output_file.name}")
        assert res["success"] is True, f"Failed on edge case: {name}"
        assert output_file.exists()

        # Clean up temporary test files
        if input_file.exists():
            input_file.unlink()
        if output_file.exists():
            output_file.unlink()

    print("=================================================================================")
    print("   ALL 5 DAY 5 EDGE CASES PASSED PERFECTLY!")
    print("=================================================================================\n")


if __name__ == "__main__":
    main()
