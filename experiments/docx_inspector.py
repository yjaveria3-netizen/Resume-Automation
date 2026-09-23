import sys
from pathlib import Path
import docx


def inspect_docx(filepath: str):
    """Inspects and prints detailed breakdown of .docx document hierarchy."""
    path = Path(filepath)
    if not path.exists():
        print(f"[ERROR] File not found: {filepath}")
        return

    doc = docx.Document(filepath)

    print("=================================================================================")
    print(f"   DOCX DOCUMENT INSPECTION REPORT: {path.name}")
    print("=================================================================================")
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}\n")

    heading_styles = set()
    bullet_styles = set()

    print("--- PARAGRAPH BREAKDOWN ---")
    for idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        style_name = p.style.name if p.style else "Normal"
        runs = p.runs
        run_count = len(runs)

        if "heading" in style_name.lower():
            heading_styles.add(style_name)
        if "bullet" in style_name.lower() or text.startswith(("•", "-", "*")):
            bullet_styles.add(style_name)

        if not text:
            print(f"[{idx:02d}] <EMPTY PARAGRAPH>")
            continue

        font_names = set()
        font_sizes = set()
        bolds = set()

        for r in runs:
            if r.font.name:
                font_names.add(r.font.name)
            if r.font.size:
                font_sizes.add(f"{r.font.size.pt}pt")
            if r.bold is not None:
                bolds.add(r.bold)

        fonts_str = ", ".join(font_names) if font_names else "Inherited"
        sizes_str = ", ".join(font_sizes) if font_sizes else "Inherited"
        bold_str = "Bold" if True in bolds else "Regular"

        preview = text[:60] + "..." if len(text) > 60 else text
        print(f"[{idx:02d}] Style: '{style_name}' | Runs: {run_count} | Font: {fonts_str} ({sizes_str}, {bold_str})")
        print(f"     Text: \"{preview}\"")

    if doc.tables:
        print("\n--- TABLE BREAKDOWN ---")
        for tidx, table in enumerate(doc.tables):
            rows = len(table.rows)
            cols = len(table.columns) if rows > 0 else 0
            print(f"Table #{tidx+1}: {rows} rows x {cols} columns")

    print("\n=================================================================================")
    print("   SUMMARY & DETECTED STYLES")
    print("=================================================================================")
    print(f"Detected Heading Styles: {', '.join(heading_styles) if heading_styles else 'None'}")
    print(f"Detected Bullet Styles:  {', '.join(bullet_styles) if bullet_styles else 'None'}")
    print("=================================================================================\n")


if __name__ == "__main__":
    if len(sys.argv) > 1:
        inspect_docx(sys.argv[1])
    else:
        print("Usage: python experiments/docx_inspector.py <path_to_sample_resume.docx>")
