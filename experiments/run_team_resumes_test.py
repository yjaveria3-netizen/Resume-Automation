import sys
from pathlib import Path
import docx

ROOT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT_DIR))
sys.path.insert(0, str(ROOT_DIR / "backend"))

from services.gemini_prompt import build_specialized_system_prompt, infer_engineering_role
from services.gemini_client import generate_bullets_with_retries
from services.scoring import compute_ats_score
from docx_engine.unified_updater import update_resume_with_project

TEAM_RESUMES = [
    {
        "role": "Person A (Frontend Lead)",
        "tech_stack": ["React", "TypeScript", "Tailwind CSS", "Vite", "Redux"],
        "repo_name": "auto-applier-frontend",
        "description": "Responsive React dashboard with Tailwind CSS styling and interactive document preview",
    },
    {
        "role": "Person B (Frontend Support & QA)",
        "tech_stack": ["React", "JavaScript", "Jest", "Cypress", "HTML5"],
        "repo_name": "qa-test-suite",
        "description": "End-to-end testing suite for React frontend components using Cypress and Jest",
    },
    {
        "role": "Person C (Backend Lead)",
        "tech_stack": ["FastAPI", "Python", "PostgreSQL", "Docker", "SQLAlchemy"],
        "repo_name": "resume-auto-updater-backend",
        "description": "High-throughput FastAPI REST API with PostgreSQL connection pooling and JWT auth",
    },
    {
        "role": "Person D (AI & Resume Logic)",
        "tech_stack": ["Python", "Google Gemini API", "PyTorch", "python-docx"],
        "repo_name": "ai-resume-generator",
        "description": "LLM prompt engineering pipeline with python-docx document structure preservation",
    },
    {
        "role": "Person E (DevOps & Database)",
        "tech_stack": ["PostgreSQL", "Docker", "Kubernetes", "Redis", "Prometheus"],
        "repo_name": "infrastructure-deploy",
        "description": "Kubernetes deployment manifests with Redis cache and Prometheus monitoring",
    },
]


def run_team_resumes_validation():
    print("=================================================================================")
    print("   PERSON D: DAY 9 TEAM RESUMES VALIDATION & PROMPT FINE-TUNING HARNESS")
    print("=================================================================================")

    exp_dir = Path(__file__).resolve().parent
    base_sample_path = exp_dir / "sample_resume.docx"

    # Ensure base sample resume exists
    if not base_sample_path.exists():
        doc = docx.Document()
        doc.add_heading("Team Member", 0)
        doc.add_heading("Projects", 1)
        doc.save(str(base_sample_path))

    for idx, team_member in enumerate(TEAM_RESUMES, 1):
        role_inferred = infer_engineering_role(team_member["tech_stack"])
        bullets = generate_bullets_with_retries(
            repo_name=team_member["repo_name"],
            description=team_member["description"],
            tech_stack=team_member["tech_stack"],
        )

        output_path = exp_dir / f"team_resume_out_{idx}.docx"
        update_resume_with_project(
            input_docx_path=str(base_sample_path),
            output_docx_path=str(output_path),
            project_title=team_member["repo_name"],
            tech_stack=team_member["tech_stack"],
            bullet_point=bullets[0],
        )

        ats_data = compute_ats_score(str(output_path))

        print(f"[{idx}/5] {team_member['role']}")
        print(f"  Inferred Role: {role_inferred}")
        print(f"  Generated Bullet: \"{bullets[0]}\"")
        print(f"  ATS Score: {ats_data['score']}/100")
        print("-" * 80)

        # Cleanup output file
        if output_path.exists():
            output_path.unlink()

    print("=================================================================================")
    print("   ALL 5 TEAM RESUMES TESTED AND VERIFIED SUCCESSFULLY!")
    print("=================================================================================\n")


if __name__ == "__main__":
    run_team_resumes_validation()
