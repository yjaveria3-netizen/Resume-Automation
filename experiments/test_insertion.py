import sys
from pathlib import Path
import docx

# Add parent root to path
ROOT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT_DIR))

from docx_engine.parser import find_projects_section_index, find_section_end_index
from docx_engine.modifier import insert_project_to_resume


def create_sample_resume(filepath: str):
    """Generates a sample resume .docx document with standard sections for testing."""
    doc = docx.Document()

    # Header Name
    p_name = doc.add_paragraph()
    r_name = p_name.add_run("Javeria Yasin")
    r_name.font.name = "Calibri"
    r_name.font.size = docx.shared.Pt(20)
    r_name.bold = True
    p_name.paragraph_format.space_after = docx.shared.Pt(6)

    # Contact Info
    p_contact = doc.add_paragraph("Email: javeria@example.com | GitHub: github.com/javeria | Location: Remote")
    p_contact.paragraph_format.space_after = docx.shared.Pt(12)

    # Technical Projects Section Header
    p_proj_head = doc.add_paragraph()
    r_proj_head = p_proj_head.add_run("Projects")
    r_proj_head.font.name = "Calibri"
    r_proj_head.font.size = docx.shared.Pt(14)
    r_proj_head.bold = True
    p_proj_head.paragraph_format.space_before = docx.shared.Pt(12)
    p_proj_head.paragraph_format.space_after = docx.shared.Pt(6)

    # Existing Project 1
    p_p1 = doc.add_paragraph()
    r_p1 = p_p1.add_run("E-Commerce Microservices Engine | Java, Spring Boot, Kafka")
    r_p1.font.name = "Calibri"
    r_p1.font.size = docx.shared.Pt(11)
    r_p1.bold = True

    p_b1 = doc.add_paragraph()
    try:
        p_b1.style = "List Bullet"
    except Exception:
        pass
    r_b1 = p_b1.add_run("Architected high-throughput order processing backend handling 5,000 requests/sec with Spring Boot.")
    r_b1.font.name = "Calibri"
    r_b1.font.size = docx.shared.Pt(10.5)

    # Experience Section Header
    p_exp_head = doc.add_paragraph()
    r_exp_head = p_exp_head.add_run("Work Experience")
    r_exp_head.font.name = "Calibri"
    r_exp_head.font.size = docx.shared.Pt(14)
    r_exp_head.bold = True
    p_exp_head.paragraph_format.space_before = docx.shared.Pt(14)
    p_exp_head.paragraph_format.space_after = docx.shared.Pt(6)

    p_exp = doc.add_paragraph("Software Engineering Intern | Tech Corp")
    p_exp_b = doc.add_paragraph("Developed automated testing suites using Python and PyTest.")

    doc.save(filepath)
    print(f"[SUCCESS] Sample resume generated at: {filepath}")


def main():
    exp_dir = Path(__file__).resolve().parent
    sample_path = exp_dir / "sample_resume.docx"
    output_path = exp_dir / "output_updated_resume.docx"

    print("=================================================================================")
    print("   PERSON D: DOCX INSERTION ENGINE TEST HARNESS (DAY 3 & DAY 4)")
    print("=================================================================================")

    # 1. Create sample resume .docx
    create_sample_resume(str(sample_path))

    # 2. Inspect original document & section parser
    doc = docx.Document(str(sample_path))
    proj_idx = find_projects_section_index(doc)
    end_idx = find_section_end_index(doc, proj_idx)

    print(f"\n[SECTION FINDER RESULT]")
    print(f"  Projects Header Index: {proj_idx} (Header Text: '{doc.paragraphs[proj_idx].text}')")
    print(f"  Projects Section End Index: {end_idx}")
    assert proj_idx != -1, "Projects section header should be detected"

    # 3. Perform insertion of new AI-generated project entry
    print(f"\n[INSERTION ENGINE]")
    print("  Inserting project: 'Resume Auto-Updater'")
    success = insert_project_to_resume(
        doc_path=str(sample_path),
        output_path=str(output_path),
        project_title="Resume Auto-Updater",
        tech_stack=["FastAPI", "React", "PostgreSQL", "Google Gemini API"],
        bullet_point="Engineered an AI-powered resume updating platform using FastAPI, React, and Google Gemini API to automate resume customization."
    )

    print(f"  Insertion Result: {'SUCCESS' if success else 'FAILED'}")
    assert success is True
    assert output_path.exists()

    # 4. Verify updated document
    doc_updated = docx.Document(str(output_path))
    print(f"\n[VERIFICATION]")
    print(f"  Original Paragraph Count: {len(doc.paragraphs)}")
    print(f"  Updated Paragraph Count:  {len(doc_updated.paragraphs)}")
    
    found_inserted = False
    for p in doc_updated.paragraphs:
        if "Resume Auto-Updater" in p.text:
            found_inserted = True
            print(f"  Found Inserted Title: \"{p.text}\"")

    assert found_inserted is True, "Inserted project title must be present in updated document"

    print("\n=================================================================================")
    print("   ALL DAY 3 AND DAY 4 DOCX ENGINE TESTS PASSED PERFECTLY!")
    print("=================================================================================\n")


if __name__ == "__main__":
    main()
