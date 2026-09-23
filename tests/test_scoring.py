import sys
import time
from pathlib import Path

# Add project root and backend to python path
ROOT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT_DIR))
sys.path.insert(0, str(ROOT_DIR / "backend"))

import docx
from services.scoring import compute_ats_score


def create_strong_resume(filepath: str):
    doc = docx.Document()
    doc.add_heading("Javeria Yasin", level=0)

    # Technical Projects
    doc.add_heading("Projects", level=1)
    p1 = doc.add_paragraph()
    p1.add_run("Resume Auto-Updater | Python, FastAPI, React, PostgreSQL, Google Gemini API, Docker").bold = True
    b1 = doc.add_paragraph(style="List Bullet")
    b1.add_run("Architected and deployed high-throughput AI resume processing pipeline using FastAPI, React, and Google Gemini API, accelerating document generation by 45%.")

    p2 = doc.add_paragraph()
    p2.add_run("Distributed Task Runner | Go, Redis, Prometheus, gRPC").bold = True
    b2 = doc.add_paragraph(style="List Bullet")
    b2.add_run("Engineered and optimized scalable task queue handling 10,000+ jobs/sec with Redis broker and Prometheus metrics monitoring.")

    # Work Experience
    doc.add_heading("Work Experience", level=1)
    doc.add_paragraph("Software Engineering Intern | Tech Corp")
    b3 = doc.add_paragraph(style="List Bullet")
    b3.add_run("Automated and streamlined CI/CD deployment pipelines using Docker, Kubernetes, and GitHub Actions, reducing deployment downtime by 30%.")

    # Skills
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Languages: Python, Go, Java, TypeScript, SQL")
    doc.add_paragraph("Frameworks & Tools: FastAPI, React, PostgreSQL, Docker, AWS, PyTorch")

    doc.save(filepath)


def create_sparse_resume(filepath: str):
    doc = docx.Document()
    doc.add_paragraph("John Doe")
    doc.add_paragraph("Student at University")
    doc.add_paragraph("Worked on simple note app.")
    doc.save(filepath)


def test_strong_resume_scores_high():
    path = Path("tests/strong_resume.docx")
    path.parent.mkdir(parents=True, exist_ok=True)
    create_strong_resume(str(path))

    start_time = time.time()
    result = compute_ats_score(str(path))
    duration_ms = (time.time() - start_time) * 1000

    print("STRONG RESUME RESULT:", result)
    assert result["score"] >= 75, f"Expected strong resume score >= 75, got {result['score']}"
    assert duration_ms < 200, f"Expected execution under 200ms, took {duration_ms}ms"

    if path.exists():
        path.unlink()


def test_sparse_resume_scores_low():
    path = Path("tests/sparse_resume.docx")
    path.parent.mkdir(parents=True, exist_ok=True)
    create_sparse_resume(str(path))

    result = compute_ats_score(str(path))
    print("SPARSE RESUME RESULT:", result)
    assert result["score"] < 50, f"Expected sparse resume score < 50, got {result['score']}"

    if path.exists():
        path.unlink()


def test_score_increments_after_bullet_added():
    sparse_path = Path("tests/sparse_before.docx")
    sparse_path.parent.mkdir(parents=True, exist_ok=True)
    create_sparse_resume(str(sparse_path))

    before_result = compute_ats_score(str(sparse_path))

    # Add strong AI bullet entry
    doc = docx.Document(str(sparse_path))
    doc.add_heading("Projects", level=1)
    p = doc.add_paragraph()
    p.add_run("AI Resume Optimizer | FastAPI, React, PostgreSQL").bold = True
    b = doc.add_paragraph(style="List Bullet")
    b.add_run("Architected high-speed backend using FastAPI, PyTorch, and PostgreSQL, increasing user engagement by 40%.")

    updated_path = Path("tests/sparse_after.docx")
    doc.save(str(updated_path))

    after_result = compute_ats_score(str(updated_path))

    print(f"BEFORE SCORE: {before_result['score']} -> AFTER SCORE: {after_result['score']}")
    assert after_result["score"] > before_result["score"], "ATS score should increase after adding strong project bullet"

    if sparse_path.exists():
        sparse_path.unlink()
    if updated_path.exists():
        updated_path.unlink()


if __name__ == "__main__":
    test_strong_resume_scores_high()
    test_sparse_resume_scores_low()
    test_score_increments_after_bullet_added()
    print("ALL ATS SCORING TESTS PASSED!")
